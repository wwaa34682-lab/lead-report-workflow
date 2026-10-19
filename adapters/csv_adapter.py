"""Read rows from an exported CSV / JSON / JSONL / XLSX file.

source:
  type: csv
  path: exports/platform_x_{date}.csv   # {date} {yyyymmdd} {month} placeholders
  encoding: utf-8-sig
  max_age_minutes: 120                   # file mtime must be fresh (live re-pull rule); null = off
"""
from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from core.config import resolve_path

from .base import FetchContext, SourceAdapter, StaleDataError


def fill_placeholders(s: str, ctx: FetchContext) -> str:
    d = ctx.data_day
    return s.format(base_dir=ctx.base_dir, date=d.isoformat(), yyyymmdd=d.strftime("%Y%m%d"), month=d.strftime("%Y-%m"),
                    period_start=ctx.period_start.isoformat())


class CsvAdapter(SourceAdapter):
    type_name = "csv"

    def fetch(self, ctx: FetchContext):
        path = resolve_path(self.global_cfg, fill_placeholders(self.cfg["path"], ctx))
        if not path.exists():
            raise FileNotFoundError(f"{self.id}: export not found: {path}")
        mtime = datetime.fromtimestamp(path.stat().st_mtime, ctx.now.tzinfo)
        max_age = self.cfg.get("max_age_minutes")
        if max_age and (ctx.now - mtime).total_seconds() > max_age * 60:
            raise StaleDataError(f"{self.id}: {path.name} is {int((ctx.now - mtime).total_seconds() // 60)} min old "
                                 f"(> {max_age}); re-export before reporting")
        return self.result(self._read(path), ctx, fetched_at=mtime)

    def _read(self, path: Path) -> list[dict]:
        enc = self.cfg.get("encoding", "utf-8-sig")
        suf = path.suffix.lower()
        if suf == ".json":
            data = json.loads(path.read_text(encoding=enc))
            for k in (self.cfg.get("records_path") or "").split(".") if self.cfg.get("records_path") else []:
                data = data[k]
            return data
        if suf == ".jsonl":
            return [json.loads(l) for l in path.read_text(encoding=enc).splitlines() if l.strip()]
        if suf in (".xlsx", ".xlsm"):
            try:
                import openpyxl  # optional dependency
            except ImportError as e:
                raise RuntimeError("pip install openpyxl to read .xlsx exports") from e
            ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[self.cfg.get("sheet") or 0] \
                if isinstance(self.cfg.get("sheet"), str) else openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[self.cfg.get("sheet", 0)]
            it = ws.iter_rows(values_only=True)
            head = [str(h).strip() if h is not None else "" for h in next(it)]
            return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]
        with path.open(encoding=enc, newline="") as f:
            return list(csv.DictReader(f, delimiter=self.cfg.get("delimiter", ",")))
